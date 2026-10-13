import json
from datetime import datetime

import openpyxl
import pytest

from api.vocab.base import VocabularyClient
from metadata_standard import MetadataStandard, Term
from processors.vocab_enhancement.lookup import VocabLookupService
from processors.vocab_enhancement.processor import VocabEnhancementProcessor
from processors.vocab_enhancement.writers import write_excel, write_json, write_report
from utils.exceptions import (VocabHttpError, WorkflowFilesNotFound,
                              WorkflowFormatError)
from utils.run_folder import RunFolder
from utils.workflow_store import Workflow, WorkflowStore


class StubClient:
    def __init__(self, name, group=None, fail_on=(), n=2):
        self.name, self.group = name, group or name
        self.fail_on, self.n = fail_on, n

    def search(self, query, subject=""):
        if query in self.fail_on:
            raise VocabHttpError("down")
        return [Term(f"{query} {i}", f"u/{self.name}/{query}/{i}", self.name)
                for i in range(self.n)]


def test_lookup_continues_after_failure_and_caps():
    svc = VocabLookupService([StubClient("LCSH", fail_on=("bad",)), StubClient("FAST"),
                              StubClient("Getty AAT", "Getty"), StubClient("Getty TGN", "Getty")],
                             [StubClient("FAST Geographic")], per_vocab_limit=3, geo_limit=1)
    r = svc.lookup(["bad", "ok"], ["Baltimore (City)"])
    assert len(r.errors) == 1 and "bad" in r.errors[0]
    assert [t.source for t in r.subject_to_terms["bad"]][0] == "FAST"
    assert len(r.subject_to_terms["ok"]) == 2 + 2 + 3   # Getty group capped at 3
    assert len(r.geographic_to_terms["Baltimore (City)"]) == 1
    assert r.geographic_to_terms["Baltimore (City)"][0].label.startswith("Baltimore 0")


def _std():
    s = MetadataStandard()
    s.add_terms("Rivers", [Term("Rivers; Lakes", "u1", "LCSH"), Term("Rivers", "u2", "FAST")])
    s.add_terms("Texas", [Term("Texas", "g1", "FAST Geographic")], geographic=True)
    return s


def _workflow():
    return Workflow(items=[{"filename": "a", "analysis": {
        "subjects": ["Rivers", "Nothing"], "geographic_entities": "Texas"}}],
        api_stats={"k": 1})


def test_write_json_keys_and_trailer(tmp_path):
    rf = RunFolder(tmp_path); rf.ensure_dirs()
    store = WorkflowStore(rf)
    n = write_json(store, _workflow(), _std())
    assert n == 1
    data = json.loads(rf.json_path.read_text())
    a = data[0]["analysis"]
    assert set(a["vocabulary_search_results"]) == {"Rivers"}
    assert a["geographic_vocabulary_search_results"]["Texas"][0]["uri"] == "g1"
    assert data[-1] == {"api_stats": {"k": 1}}


def _make_xlsx(rf, headers=("Filename", "Subjects", "Geographic Entities", "Content Warning")):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Analysis"
    ws.append(list(headers))
    ws.append(["a", "Rivers, Nothing", "Texas", "None"])
    wb.save(rf.xlsx_path)


def test_write_excel_by_header_and_semicolon_label(tmp_path):
    rf = RunFolder(tmp_path); rf.ensure_dirs(); _make_xlsx(rf)
    write_excel(WorkflowStore(rf), _std())
    ws = openpyxl.load_workbook(rf.xlsx_path)["Analysis"]
    assert ws.cell(2, 5).value == "Rivers; Lakes (u1) [LCSH]; Rivers (u2) [FAST]"
    assert ws.cell(2, 6).value == "Texas (g1) [FAST Geographic]"


def test_write_excel_missing_header(tmp_path):
    rf = RunFolder(tmp_path); rf.ensure_dirs()
    _make_xlsx(rf, headers=("Filename", "Subjects", "Content Warning"))
    with pytest.raises(WorkflowFormatError):
        write_excel(WorkflowStore(rf), _std())


def test_report_precedence_and_counts(tmp_path):
    p = tmp_path / "r.txt"
    write_report(p, _std(), workflow_type="", clock=lambda: datetime(2026, 1, 1))
    t = p.read_text()
    assert "Workflow Type: UNKNOWN" in t
    assert "- LCSH: 1 terms" in t and "- FAST Geographic: 1 terms" in t


def test_processor_end_to_end_and_errors(tmp_path):
    rf = RunFolder(tmp_path)
    svc = VocabLookupService([StubClient("LCSH")], [StubClient("FAST Geographic")])
    proc = VocabEnhancementProcessor(rf, svc)
    with pytest.raises(WorkflowFilesNotFound):
        proc.run()
    rf.ensure_dirs(); _make_xlsx(rf)
    WorkflowStore(rf).save(_workflow())
    s = proc.run()
    assert s.n_subjects == 2 and s.subjects_with_terms == 2 and not s.errors
    assert rf.vocab_report_path.name == "vocabulary_mapping_report.txt"
    assert rf.vocab_report_path.exists()


def test_processor_no_subjects(tmp_path):
    rf = RunFolder(tmp_path); rf.ensure_dirs(); _make_xlsx(rf)
    WorkflowStore(rf).save(Workflow(items=[{"analysis": {}}]))
    s = VocabEnhancementProcessor(rf, VocabLookupService([], [])).run()
    assert s.nothing_to_do
