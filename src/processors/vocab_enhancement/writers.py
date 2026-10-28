"""Output writers: JSON, Excel and the text reports."""
from __future__ import annotations

import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Mapping, Union

from openpyxl.styles import Alignment

from metadata_standard import MetadataStandard
from utils.exceptions import WorkflowFormatError
from utils.logging_setup import run_parameters_block
from utils.subjects import parse_subject_list
from utils.workflow_store import Workflow, WorkflowStore

from .extraction import extract_subjects_and_geo

SHEET = "Analysis"
SUBJECT_VOCAB_HEADER = "Subject Vocabulary Terms (Max 3 per vocab per subject)"
GEO_VOCAB_HEADER = "Geographic Vocabulary Terms (FAST only)"


def write_json(store: WorkflowStore, workflow: Workflow,
               standards: MetadataStandard) -> int:
    """Add per-item vocabulary mappings to the workflow and save it.
    Returns the number of items that received terms."""
    subj_json, geo_json = standards.to_json(), standards.geographic_to_json()
    processed = 0
    for item in workflow.items:
        analysis = item.get("analysis")
        if not isinstance(analysis, dict):
            continue
        subjects, geos = extract_subjects_and_geo(item)
        s = {k: copy.deepcopy(subj_json[k]) for k in subjects if subj_json.get(k)}
        g = {k: copy.deepcopy(geo_json[k]) for k in geos if geo_json.get(k)}
        analysis["vocabulary_search_results"] = s
        analysis["geographic_vocabulary_search_results"] = g
        if s or g:
            processed += 1
    store.save(workflow)
    return processed


def _cell_text(keys: list[str], strings: Mapping[str, list[str]]) -> str:
    seen: list[str] = []
    for k in keys:
        for entry in strings.get(k, []):
            if entry not in seen:
                seen.append(entry)
    return "; ".join(seen)


def write_excel(store: WorkflowStore, standards: MetadataStandard) -> int:
    """Insert vocabulary columns after 'Content Warning', locating columns
    by header. Raises WorkflowFormatError when a header is missing."""
    wb = store.open_workbook()
    if SHEET not in wb.sheetnames:
        raise WorkflowFormatError(f"Sheet '{SHEET}' not found in workbook")
    ws = wb[SHEET]
    subj_col = store.column_index(ws, "Subjects")
    geo_col = store.column_index(ws, "Geographic Entities")
    insert_col = store.column_index(ws, "Content Warning") + 1
    # insert_cols shifts later columns right; read inputs after insert.
    ws.insert_cols(insert_col, 2)
    if subj_col >= insert_col:
        subj_col += 2
    if geo_col >= insert_col:
        geo_col += 2
    for col, header in ((insert_col, SUBJECT_VOCAB_HEADER),
                        (insert_col + 1, GEO_VOCAB_HEADER)):
        c = ws.cell(row=1, column=col)
        c.value = header
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 50
    s_strings = {k: [t.format() for t in ts if t.label]
                 for k, ts in standards.subject_to_terms.items()}
    g_strings = {k: [t.format() for t in ts if t.label]
                 for k, ts in standards.geographic_to_terms.items()}
    processed = 0
    for row in range(2, ws.max_row + 1):
        subjects = parse_subject_list(ws.cell(row=row, column=subj_col).value)
        geos = parse_subject_list(ws.cell(row=row, column=geo_col).value)
        s_text, g_text = _cell_text(subjects, s_strings), _cell_text(geos, g_strings)
        for col, text in ((insert_col, s_text), (insert_col + 1, g_text)):
            c = ws.cell(row=row, column=col)
            c.value = text
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if s_text or g_text:
            processed += 1
    store.save_workbook(wb)
    return processed


def write_report(path: Union[str, Path], standards: MetadataStandard, *,
                 workflow_type: str = "text", max_terms_per_vocab: int = 3,
                 max_total_terms: int = 12,
                 clock: Callable[[], datetime] = datetime.now) -> None:
    subj, geo = standards.subject_to_terms, standards.geographic_to_terms
    s_with = sum(1 for t in subj.values() if t)
    g_with = sum(1 for t in geo.values() if t)

    def counts(mapping: Mapping[str, list], sources: list[str]) -> dict[str, int]:
        c = {s: 0 for s in sources}
        for ts in mapping.values():
            for t in ts:
                c[t.source] = c.get(t.source, 0) + 1
        return c

    def rate(n: int, d: int) -> str:
        return f"{n / d * 100:.1f}%" if d else "0%"

    L = ["DISSERTATION VOCABULARY MAPPING REPORT", "=" * 50, "",
         f"Generated: {clock():%Y-%m-%d %H:%M:%S}",
         f"Workflow Type: {(workflow_type or 'Unknown').upper()}",
         f"Total Subjects Processed: {len(subj)}",
         f"Total Geographic Entities Processed: {len(geo)}",
         f"Terms Per Vocabulary Limit: {max_terms_per_vocab}",
         f"Max Total Terms Per Subject: {max_total_terms}", "",
         "VOCABULARIES SEARCHED:", "FOR SUBJECTS:",
         "- LCSH (Library of Congress Subject Headings)",
         "- FAST (Faceted Application of Subject Terminology)",
         "- Getty AAT (Art & Architecture Thesaurus)",
         "- Getty TGN (Thesaurus of Geographic Names)",
         "FOR GEOGRAPHIC ENTITIES:",
         "- FAST Geographic (Faceted Application of Subject Terminology - Geographic)",
         "", "STATISTICS:",
         f"- Subjects with vocabulary terms: {s_with}",
         f"- Subjects without vocabulary terms: {len(subj) - s_with}",
         f"- Subject success rate: {rate(s_with, len(subj))}",
         f"- Geographic entities with vocabulary terms: {g_with}",
         f"- Geographic entities without vocabulary terms: {len(geo) - g_with}",
         f"- Geographic success rate: {rate(g_with, len(geo))}", "",
         "SUBJECT TERMS BY SOURCE:"]
    L += [f"- {s}: {n} terms" for s, n in counts(
        subj, ["LCSH", "FAST", "Getty AAT", "Getty TGN"]).items()]
    L += ["", "GEOGRAPHIC TERMS BY SOURCE:"]
    L += [f"- {s}: {n} terms" for s, n in counts(geo, ["FAST Geographic"]).items()]
    L += ["", "ALPHABETICAL SUBJECT REFERENCE:", "-" * 40, ""]
    strings = standards.to_excel_strings()
    for s in sorted(subj):
        L += [f"Subject: {s}", f"Vocabulary Terms: {strings.get(s) or 'No terms found'}", ""]
    L += ["ALPHABETICAL GEOGRAPHIC ENTITY REFERENCE:", "-" * 45, ""]
    gstrings = standards.to_excel_strings(geographic=True)
    for e in sorted(geo):
        L += [f"Geographic Entity: {e}",
              f"Vocabulary Terms: {gstrings.get(e) or 'No terms found'}"]
        if any(t.source == "FAST Geographic" for t in geo[e]):
            L.append("Note: FAST Geographic terms may include additional "
                     "metadata in detailed API logs.")
        L.append("")
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")


def write_usage_log(path: Union[str, Path], stats: dict, total_subjects: int,
                    clock: Callable[[], datetime] = datetime.now) -> None:
    L = ["DISSERTATION VOCABULARY API USAGE LOG", "=" * 50, "",
         f"Generated: {clock():%Y-%m-%d %H:%M:%S}", "",
         *run_parameters_block(50).rstrip("\n").split("\n"), "", "SUMMARY STATISTICS:",
         "-" * 20, f"Total subjects processed: {total_subjects}",
         f"Total API requests made: {stats.get('total_requests', 0)}",
         f"Total cache hits: {stats.get('total_cache_hits', 0)}",
         f"Total processing time: {stats.get('total_time', 0):.2f} seconds",
         f"Average time per subject: {stats.get('avg_time_per_subject', 0):.2f} seconds", ""]
    for name, s in stats.get("api_breakdown", {}).items():
        L += [f"{name.upper()} API STATISTICS:", "-" * (len(name) + 16),
              f"Requests made: {s['requests']}",
              f"Successful responses: {s['successful']}",
              f"Failed responses: {s['failed']}",
              f"Success rate: {s['success_rate']:.1f}%",
              f"Total processing time: {s['total_time']:.2f} seconds",
              f"Average time per request: {s['avg_time']:.2f} seconds",
              f"Terms found: {s['terms_found']}",
              f"Cache hits: {s['cache_hits']}", ""]
    L += ["SUBJECT PROCESSING RESULTS:", "-" * 25]
    for subject, r in stats.get("subject_results", {}).items():
        L += [f"Subject: {subject}", f"  Total terms found: {r['total_terms']}"]
        L += [f"  {a}: {n} terms" for a, n in r["terms_by_api"].items()]
        L += [f"  Processing time: {r['processing_time']:.2f}s", ""]
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")
