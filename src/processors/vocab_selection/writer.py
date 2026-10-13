"""Writes selection results to JSON, Excel and a text report."""
from __future__ import annotations

import copy
from datetime import datetime
from typing import Callable, Mapping

from openpyxl.styles import Alignment

from metadata_standard import MetadataStandard, Term
from utils.workflow_store import Workflow, WorkflowStore

from .selector import SelectionResult, SelectionStatus

HEADER = "Selected Subject Vocabulary Terms"


def term_to_json(t: Term) -> dict:
    d = {"label": t.label, "uri": t.uri, "source": t.source}
    for k, v in t.to_dict().items():
        if k not in d and v:
            d[k] = v
    return d


class SelectionWriter:
    def __init__(self, store: WorkflowStore,
                 clock: Callable[[], datetime] = datetime.now) -> None:
        self.store = store
        self.clock = clock

    def apply_to_workflow(self, workflow: Workflow,
                          results: Mapping[int, SelectionResult]) -> Workflow:
        """Returns a new Workflow; the input is not mutated."""
        items = copy.deepcopy(workflow.items)
        for i, item in enumerate(items):
            analysis = item.setdefault("analysis", {})
            r = results.get(i)
            analysis["final_selected_terms"] = [term_to_json(t) for t in r.terms] if r else []
            if r is not None:
                analysis["selection_status"] = r.status.value
                if r.error:
                    analysis["selection_error"] = r.error
                if r.unmatched_labels:
                    analysis["selection_unmatched_labels"] = list(r.unmatched_labels)
        return Workflow(items=items, api_stats=copy.deepcopy(workflow.api_stats))

    def write_json(self, workflow: Workflow, results: Mapping[int, SelectionResult]) -> Workflow:
        new = self.apply_to_workflow(workflow, results)
        self.store.save(new)
        return new

    def write_excel(self, workflow: Workflow) -> None:
        """workflow: result of apply_to_workflow (reads final_selected_terms)."""
        wb = self.store.open_workbook()
        ws = wb["Analysis"]
        col = self.store.column_index(ws, "Subject Vocabulary Terms")
        ws.cell(row=1, column=col).value = HEADER
        for i, item in enumerate(workflow.items):
            terms = item.get("analysis", {}).get("final_selected_terms", [])
            text = "; ".join(
                Term(t.get("label", ""), t.get("uri", ""), t.get("source", "")).format()
                for t in terms if t.get("label"))
            cell = ws.cell(row=i + 2, column=col)
            cell.value = text
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in range(len(workflow.items) + 2, ws.max_row + 1):
            ws.cell(row=row, column=col).value = ""
        self.store.save_workbook(wb)

    def write_selection_report(self, workflow: Workflow,
                               results: Mapping[int, SelectionResult]) -> dict:
        """Write the report; returns summary counts."""
        failed = sum(1 for r in results.values() if r.status is SelectionStatus.FAILED)
        unmatched = sum(1 for r in results.values() if r.unmatched_labels)
        issues = sum(1 for r in results.values() if r.has_issue)
        selected = sum(1 for r in results.values() if r.status is SelectionStatus.SELECTED)
        lines = ["DISSERTATION VOCABULARY SELECTION REPORT (STEP 3)", "=" * 60, "",
                 f"Generated: {self.clock():%Y-%m-%d %H:%M:%S}",
                 f"Entries: {len(workflow.items)}  Selected: {selected}  "
                 f"Failed: {failed}  Items with issues: {issues}", ""]
        for i, item in enumerate(workflow.items):
            a = item.get("analysis", {})
            r = results.get(i)
            lines += [f"ENTRY {i}:", f"Title: {a.get('title', 'Unknown title')}",
                      f"Status: {r.status.value if r else 'not_processed'}"]
            if r and r.error:
                lines.append(f"Error: {r.error}")
            if r and r.unmatched_labels:
                lines.append("Unmatched labels: " + "; ".join(r.unmatched_labels))
            lines.append("-" * 60)
            terms = a.get("final_selected_terms", [])
            if terms:
                lines.append("SELECTED CONTROLLED TERMS (Step 3):")
                lines += [f"  - {Term(t['label'], t.get('uri', ''), t.get('source', '')).format()}"
                          for t in terms]
            else:
                lines.append("No controlled terms selected for this entry.")
            lines += ["", "=" * 60, ""]
        path = self.store.run_folder.selection_report_path
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("\n".join(lines), encoding="utf-8")
        return {"total": len(results), "selected": selected, "failed": failed,
                "unmatched": unmatched, "items_with_issues": issues,
                "report_path": str(path)}
