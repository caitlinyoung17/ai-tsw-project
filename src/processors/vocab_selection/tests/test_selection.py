import json

import openpyxl
import pytest

from api.llm_core.testing import FakeProvider
from api.llm_core.types import BatchOutcome, CompletionResult, Usage
from metadata_standard import MetadataStandard, Term
from processors.vocab_selection.batch import SelectionBatchRunner
from processors.vocab_selection.processor import Step2NotRunError, VocabSelectionProcessor
from processors.vocab_selection.selector import (ExtractedEntry, SelectionStatus,
                                                 VocabularySelector)
from processors.vocab_selection.writer import SelectionWriter
from utils.exceptions import LLMError, LLMTransientError, WorkflowFilesNotFound
from utils.run_folder import RunFolder
from utils.workflow_store import Workflow, WorkflowStore

T = lambda l, s, u="": Term(l, u, s)
STD = MetadataStandard({"art": [T("Art history", "LCSH", "http://a"),
                                T("Art history", "FAST", "http://b"),
                                T("Painting", "Getty AAT", "http://c")]})
E = ExtractedEntry(0, "Title", "Abstract")
sleeps = []


def sel(provider, **kw):
    return VocabularySelector(provider, "fake-model", sleep=sleeps.append, **kw)


def good(*labels):
    return json.dumps({"selected_terms": [{"label": l} for l in labels]})


def test_render_matches_legacy_format():
    std = MetadataStandard({"a": [Term("X", "http://u", "LCSH"), Term("Y", "", "FAST")],
                            "b": []})
    assert std.render_for_prompt() == (
        "Subject: a\n  - X (http://u) [LCSH]\n  - Y [FAST]\n\n"
        "Subject: b\n  (No candidate controlled terms)\n")


def test_skipped_no_candidates():
    p = FakeProvider([])
    r = sel(p).select(E, MetadataStandard({"x": []}))
    assert r.status is SelectionStatus.SKIPPED_NO_CANDIDATES and p.requests == []


def test_selected_and_priority_and_unmatched():
    p = FakeProvider([good("art history", "Painting", "Nonexistent")])
    r = sel(p).select(E, STD)
    assert r.status is SelectionStatus.SELECTED
    assert [t.label for t in r.terms] == ["Art history", "Painting"]
    assert r.terms[0].source == "LCSH"
    assert r.unmatched_labels == ["Nonexistent"] and r.usage.total == 15
    assert p.requests[0].max_tokens == 1500


def test_none_selected():
    r = sel(FakeProvider(['{"selected_terms": []}'])).select(E, STD)
    assert r.status is SelectionStatus.NONE_SELECTED and r.error is None


def test_malformed_json_failed():
    r = sel(FakeProvider(["not json"])).select(E, STD)
    assert r.status is SelectionStatus.FAILED and r.error and r.terms == []
    r = sel(FakeProvider(['{"foo": 1}'])).select(E, STD)
    assert r.status is SelectionStatus.FAILED


def test_transient_retried_nontransient_not():
    p = FakeProvider([LLMTransientError("429"), good("Painting")])
    assert sel(p).select(E, STD).status is SelectionStatus.SELECTED
    assert len(p.requests) == 2
    p = FakeProvider([LLMError("bad"), good("Painting")])
    r = sel(p).select(E, STD)
    assert r.status is SelectionStatus.FAILED and len(p.requests) == 1


class BatchFake(FakeProvider):
    supports_batch = True

    def __init__(self, outcomes):
        super().__init__([])
        self.outcomes = outcomes
        self.submitted = []

    def submit_batch(self, items):
        self.submitted = items
        return "h"

    def collect_batch(self, handle, timeout=None, poll_interval=60.0, sleep=None):
        return {k: v(k) if callable(v) else v for k, v in self.outcomes.items()}


def ok(text):
    return BatchOutcome("x", CompletionResult(text, Usage(1, 1), "m"))


def test_batch_mapping_skip_and_mixed():
    entries = [(ExtractedEntry(0), STD), (ExtractedEntry(1), MetadataStandard()),
               (ExtractedEntry(2), STD), (ExtractedEntry(3), STD)]
    p = BatchFake({"entry-0": ok(good("Painting")),
                   "entry-2": BatchOutcome("entry-2", error="boom")})
    res = SelectionBatchRunner(p, sel(p)).run(entries)
    assert [i.custom_id for i in p.submitted] == ["entry-0", "entry-2", "entry-3"]
    assert res[0].status is SelectionStatus.SELECTED
    assert res[1].status is SelectionStatus.SKIPPED_NO_CANDIDATES
    assert res[2].status is SelectionStatus.FAILED and res[2].error == "boom"
    assert res[3].status is SelectionStatus.FAILED  # missing outcome


def make_run(tmp_path, with_vocab=True):
    rf = RunFolder(tmp_path / "run")
    rf.ensure_dirs()
    a = {"title": "T0", "abstract": "A", "subjects": ["art"]}
    if with_vocab:
        a["vocabulary_search_results"] = {"art": [t.to_dict() for t in STD.terms_for("art")]}
    items = [{"analysis": a}, {"analysis": {"title": "T1",
             **({"vocabulary_search_results": {}} if with_vocab else {})}}]
    store = WorkflowStore(rf)
    store.save(Workflow(items, {"total_requests": 1}))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Analysis"
    ws.append(["Title", "Subject Vocabulary Terms"])
    ws.append(["T0", "old"])
    ws.append(["T1", "old"])
    wb.save(rf.xlsx_path)
    return rf, store


def test_writer_and_processor(tmp_path):
    rf, store = make_run(tmp_path)
    p = FakeProvider([good("Painting", "Zzz")])
    summary = VocabSelectionProcessor(rf, p, "fake-model").run()
    assert summary["items_with_issues"] == 1 and summary["selected"] == 1
    data = json.loads(rf.json_path.read_text())
    assert data[0]["analysis"]["final_selected_terms"][0]["label"] == "Painting"
    assert data[0]["analysis"]["selection_status"] == "selected"
    assert data[1]["analysis"]["final_selected_terms"] == []
    assert "api_stats" in data[-1]
    ws = openpyxl.load_workbook(rf.xlsx_path)["Analysis"]
    assert ws["B1"].value == "Selected Subject Vocabulary Terms"
    assert ws["B2"].value == "Painting (http://c) [Getty AAT]" and not ws["B3"].value
    assert "Unmatched labels: Zzz" in rf.selection_report_path.read_text()


def test_writer_does_not_mutate(tmp_path):
    rf, store = make_run(tmp_path)
    wf = store.load()
    r = sel(FakeProvider([good("Painting")])).select(E, STD)
    new = SelectionWriter(store).apply_to_workflow(wf, {0: r})
    assert "final_selected_terms" not in wf.items[0]["analysis"]
    assert new.items[0]["analysis"]["final_selected_terms"]


def test_failed_counts_in_report(tmp_path):
    rf, _ = make_run(tmp_path)
    s = VocabSelectionProcessor(rf, FakeProvider(["junk"]), "m").run()
    assert s["failed"] == 1 and s["items_with_issues"] == 1
    d = json.loads(rf.json_path.read_text())
    assert d[0]["analysis"]["selection_status"] == "failed" and d[0]["analysis"]["selection_error"]


def test_errors(tmp_path):
    rf, _ = make_run(tmp_path, with_vocab=False)
    with pytest.raises(Step2NotRunError):
        VocabSelectionProcessor(rf, FakeProvider([]), "m").run()
    with pytest.raises(WorkflowFilesNotFound):
        VocabSelectionProcessor(RunFolder(tmp_path / "none"), FakeProvider([]), "m").run()
