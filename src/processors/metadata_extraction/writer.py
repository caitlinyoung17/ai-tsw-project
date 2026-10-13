"""Build text_workflow.xlsx / .json once from outcomes (single path for both modes)."""
from __future__ import annotations

from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment

from utils import schema
from utils.workflow_store import Workflow, WorkflowStore

from .collector import Document, SkippedFile
from .results import ExtractionOutcome, RunSummary, Succeeded

ANALYSIS_HEADERS = ["Filename", "Title", "Creator", "Date", "Department", "Abstract",
                    "Subjects", "Type", "Access", "Publisher", "Sponsor", "Language",
                    "Named Entities", "Geographic Entities", "Content Warning"]
ANALYSIS_WIDTHS = [30, 50, 30, 12, 30, 60, 40, 15, 15, 25, 25, 15, 40, 40, 30]
_WRAP = Alignment(vertical="top", wrap_text=True)


def analysis_row(filename: str, m: dict) -> list:
    g = m.get
    return [filename, g(schema.TITLE, ""), g(schema.CREATOR, ""), g(schema.DATE, ""),
            g(schema.DEPARTMENT, ""), g(schema.ABSTRACT, ""),
            ", ".join(g(schema.SUBJECTS, [])), g(schema.TYPE, ""),
            g(schema.ACCESS, ""), g(schema.PUBLISHER, ""), g(schema.SPONSOR, ""),
            g(schema.LANGUAGE, ""), ", ".join(g(schema.NAMED_ENTITIES, [])),
            ", ".join(g(schema.GEOGRAPHIC_ENTITIES, [])),
            g(schema.CONTENT_WARNING, "None")]


def json_item(document: Document, metadata: dict, raw: str) -> dict:
    analysis = schema.to_snake(metadata)
    analysis["raw_response"] = raw
    return {"filename": document.filename, "file_path": document.path,
            "analysis": analysis}


class WorkflowWriter:
    def __init__(self, store: WorkflowStore) -> None:
        self.store = store

    def build_items(self, documents: Sequence[Document],
                    outcomes: Dict[str, ExtractionOutcome]) -> List[dict]:
        items = []
        for doc in documents:
            o = outcomes.get(doc.doc_id)
            if isinstance(o, Succeeded):
                items.append(json_item(doc, o.result.metadata, o.result.raw))
            else:
                items.append(json_item(doc, schema.empty_metadata(), ""))
        return items

    def api_stats(self, summary: RunSummary) -> dict:
        return {"total_requests": summary.total,
                "total_input_tokens": summary.usage.prompt_tokens,
                "total_output_tokens": summary.usage.completion_tokens,
                "total_tokens": summary.usage.total,
                "processing_mode": summary.mode}

    def build_workbook(self, documents, outcomes, summary: RunSummary,
                       skipped: Sequence[SkippedFile] = ()) -> Workbook:
        wb = Workbook()
        analysis = wb.active
        analysis.title = "Analysis"
        analysis.append(ANALYSIS_HEADERS)
        for i, w in enumerate(ANALYSIS_WIDTHS, start=1):
            analysis.column_dimensions[analysis.cell(row=1, column=i).column_letter].width = w
        raw_sheet = wb.create_sheet("Raw Responses")
        raw_sheet.append(["Filename", "API Response"])
        issues = wb.create_sheet("Issues")
        issues.append(["Filename", "Error"])
        for doc in documents:
            o = outcomes.get(doc.doc_id)
            if isinstance(o, Succeeded):
                meta, raw = o.result.metadata, o.result.raw
            else:
                meta, raw = schema.empty_metadata(), ""
                issues.append([doc.filename, str(o.error) if o else "No result"])
            analysis.append(analysis_row(doc.filename, meta))
            raw_sheet.append([doc.filename, raw])
            for sheet in (analysis, raw_sheet):
                for cell in sheet[sheet.max_row]:
                    cell.alignment = _WRAP
        for s in skipped:
            issues.append([s.path, s.reason])
        stats = wb.create_sheet("API Stats")
        stats.append(["Metric", "Value"])
        for k, v in self.api_stats(summary).items():
            stats.append([k, v])
        return wb

    def write(self, documents, outcomes, summary: RunSummary,
              skipped: Sequence[SkippedFile] = ()) -> None:
        wb = self.build_workbook(documents, outcomes, summary, skipped)
        self.store.run_folder.ensure_dirs()
        self.store.save_workbook(wb)
        self.store.save(Workflow(items=self.build_items(documents, outcomes),
                                 api_stats=self.api_stats(summary)))
