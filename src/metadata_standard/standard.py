"""MetadataStandard: per-run candidate controlled-vocabulary terms."""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

import openpyxl

from utils.exceptions import (PayloadValidationError, VocabularyFormatError,
                              VocabularyLoadError)
from utils.text import normalize_for_comparison

from .term import Term

# "label (uri) [source]" or "label [source]". Entries are split on ';' only
# when followed by a new entry start; a ';' inside a label is tolerated only
# if the text after it does not itself look like "... [source]" on its own.
# Limitation: labels containing "[...]" or "(http...)" text may mis-parse.
_ENTRY_RE = re.compile(
    r"\s*(?P<label>[^\[\]]*?)\s*(?:\((?P<uri>[^()\s]*)\)\s*)?\[(?P<source>[^\[\]]+)\]\s*(?:;|$)")


def _parse_terms_dict(data: Any, what: str) -> dict[str, list[Term]]:
    if not isinstance(data, Mapping):
        raise VocabularyFormatError(
            f"{what} must be a mapping of subject -> list of terms, got {type(data).__name__}")
    out: dict[str, list[Term]] = {}
    for subject, terms in data.items():
        if not isinstance(terms, list):
            raise VocabularyFormatError(f"{what}[{subject!r}] must be a list of term dicts")
        out[str(subject)] = [Term.from_dict(t) for t in terms]
    return out


def parse_excel_terms(cell: str) -> list[Term]:
    """Parse a '; '-joined 'label (uri) [source]' cell into Terms."""
    terms = []
    for m in _ENTRY_RE.finditer(cell or ""):
        label = m.group("label").strip()
        if label:
            terms.append(Term(label=label, uri=m.group("uri") or "",
                              source=m.group("source").strip()))
    return terms


class MetadataStandard:
    def __init__(self,
                 subject_to_terms: Optional[dict[str, list[Term]]] = None,
                 geographic_to_terms: Optional[dict[str, list[Term]]] = None) -> None:
        self.subject_to_terms: dict[str, list[Term]] = dict(subject_to_terms or {})
        self.geographic_to_terms: dict[str, list[Term]] = dict(geographic_to_terms or {})

    # ---- mutation / queries ----
    def add_terms(self, subject: str, terms: Iterable[Term], geographic: bool = False) -> None:
        target = self.geographic_to_terms if geographic else self.subject_to_terms
        target.setdefault(subject, []).extend(terms)

    def terms_for(self, subject: str) -> list[Term]:
        """Terms for a subject (subject or geographic); [] if unknown."""
        return list(self.subject_to_terms.get(subject, [])
                    or self.geographic_to_terms.get(subject, []))

    def has_candidates(self) -> bool:
        """True if any subject has at least one candidate term."""
        return any(self.subject_to_terms.values())

    # ---- constructors ----
    @classmethod
    def from_dict(cls, data: Mapping[str, Any],
                  geographic: Optional[Mapping[str, Any]] = None) -> "MetadataStandard":
        return cls(_parse_terms_dict(data, "vocabulary results"),
                   _parse_terms_dict(geographic, "geographic vocabulary results")
                   if geographic is not None else {})

    @classmethod
    def load_from_json(cls, path) -> "MetadataStandard":
        p = Path(path)
        try:
            raw = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            raise VocabularyLoadError(f"Cannot read {p}", path=str(p), cause=e)
        if isinstance(raw, Mapping) and "subject_to_terms" in raw:
            return cls.from_dict(raw["subject_to_terms"], raw.get("geographic_to_terms"))
        return cls.from_dict(raw)

    @classmethod
    def load_from_excel(cls, path, sheet: str = "Analysis") -> "MetadataStandard":
        """Fallback loader. Excel cells are flattened (the per-subject mapping
        is lost), so terms are keyed by each row's Subjects / Geographic
        Entities cell text. A ';' inside a label is tolerated, but labels
        containing '[...]' or '(...)' fragments may mis-parse."""
        p = Path(path)
        try:
            wb = openpyxl.load_workbook(p, read_only=False)
        except Exception as e:
            raise VocabularyLoadError(f"Cannot read {p}", path=str(p), cause=e)
        if sheet not in wb.sheetnames:
            raise VocabularyFormatError(f"Sheet '{sheet}' not found in {p}", path=str(p))
        ws = wb[sheet]
        headers = {str(c.value).lower(): c.column for c in ws[1] if c.value is not None}

        def col(*needles: str) -> Optional[int]:
            for h, idx in headers.items():
                if all(n in h for n in needles):
                    return idx
            return None

        s_key, s_voc = col("subjects"), col("subject vocabulary")
        g_key, g_voc = col("geographic entities"), col("geographic vocabulary")
        if s_voc is None and g_voc is None:
            raise VocabularyFormatError(f"No vocabulary header found in {p}", path=str(p))
        std = cls()
        for row in range(2, ws.max_row + 1):
            for key_c, voc_c, geo in ((s_key, s_voc, False), (g_key, g_voc, True)):
                if key_c is None or voc_c is None:
                    continue
                key = str(ws.cell(row=row, column=key_c).value or "").strip()
                terms = parse_excel_terms(str(ws.cell(row=row, column=voc_c).value or ""))
                if key and terms:
                    std.add_terms(key, terms, geographic=geo)
        return std

    @classmethod
    def from_workflow_item(cls, item: Mapping[str, Any]) -> "MetadataStandard":
        analysis = item.get("analysis") if isinstance(item, Mapping) else None
        if not isinstance(analysis, Mapping):
            raise VocabularyFormatError("Workflow item has no 'analysis' mapping")
        return cls.from_dict(analysis.get("vocabulary_search_results") or {},
                             analysis.get("geographic_vocabulary_search_results") or {})

    # ---- serialization ----
    def to_json(self) -> dict[str, list[dict[str, str]]]:
        return {s: [t.to_dict() for t in ts] for s, ts in self.subject_to_terms.items()}

    def geographic_to_json(self) -> dict[str, list[dict[str, str]]]:
        return {s: [t.to_dict() for t in ts] for s, ts in self.geographic_to_terms.items()}

    def to_excel_strings(self, geographic: bool = False) -> dict[str, str]:
        src = self.geographic_to_terms if geographic else self.subject_to_terms
        return {s: "; ".join(t.format() for t in ts if t.label) for s, ts in src.items()}

    # ---- prompt rendering ----
    def _build_subject_organized_terms(self) -> list[tuple[str, list[Term]]]:
        """Organized structure: [(subject, terms), ...] in insertion order."""
        return list(self.subject_to_terms.items())

    def render_for_prompt(self) -> str:
        if not self.has_candidates():
            return ""
        lines: list[str] = []
        for subject, terms in self._build_subject_organized_terms():
            lines.append(f"Subject: {subject}")
            if not terms:
                lines.append("  (No candidate controlled terms)")
            else:
                lines.extend(f"  - {t.format()}" for t in terms)
            lines.append("")
        return "\n".join(lines)

    # ---- selection checks ----
    @staticmethod
    def validate_selection(payload: Mapping[str, Any]) -> list[str]:
        """Return the selected labels; raise PayloadValidationError on bad shape.
        An empty list is valid."""
        if not isinstance(payload, Mapping) or "selected_terms" not in payload:
            raise PayloadValidationError("Payload missing 'selected_terms'")
        sel = payload["selected_terms"]
        if not isinstance(sel, list):
            raise PayloadValidationError("'selected_terms' must be a list")
        labels: list[str] = []
        for item in sel:
            if isinstance(item, str):
                labels.append(item)
            elif isinstance(item, Mapping) and isinstance(item.get("label"), str):
                labels.append(item["label"])
            else:
                raise PayloadValidationError(
                    f"'selected_terms' entries must be strings or dicts with 'label': {item!r}")
        return labels

    def _pool(self, subject: Optional[str]) -> list[Term]:
        if subject is not None:
            return list(self.subject_to_terms.get(subject, []))
        return [t for ts in self.subject_to_terms.values() for t in ts]

    @staticmethod
    def _best_for(label: str, pool: list[Term]) -> Optional[Term]:
        words = set(normalize_for_comparison(label).split())
        cands = [t for t in pool
                 if words and set(normalize_for_comparison(t.label).split()) == words]
        if not cands:
            low = label.lower().strip()
            cands = [t for t in pool if t.label.lower().strip() == low]
        return min(cands, key=lambda t: t.priority) if cands else None

    def match_selected_labels(self, selected_labels: Iterable[str],
                              subject: Optional[str] = None) -> list[Term]:
        pool = self._pool(subject)
        matched = [b for b in (self._best_for(l, pool) for l in selected_labels) if b]
        return self.dedupe(matched)

    def unmatched_labels(self, selected_labels: Iterable[str],
                         subject: Optional[str] = None) -> list[str]:
        pool = self._pool(subject)
        return [l for l in selected_labels if self._best_for(l, pool) is None]

    @staticmethod
    def dedupe(terms: Iterable[Term]) -> list[Term]:
        """Dedupe by URI keeping first occurrence; URI-less terms are kept."""
        seen: set[str] = set()
        out: list[Term] = []
        for t in terms:
            if t.uri:
                if t.uri in seen:
                    continue
                seen.add(t.uri)
            out.append(t)
        return out
